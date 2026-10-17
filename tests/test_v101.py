"""v1.0.1 regression tests; temporary databases, no real notifications.
Run: python -m unittest tests.test_v101 -v
"""
import io
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
import httpx
from openpyxl import Workbook, load_workbook
import yaml
from fastapi import FastAPI
from firefly import __version__, config as settings
from firefly.config import Config
from firefly.db import session as db
from firefly.api.router import api_router
from firefly.services import excel_import, updater
from firefly.render import bot

HEADERS=['交易日期','收支类型','金额','类别','账本名称','交易对方','备注','是否有发票']
ROW=['2026-10-17 12:30','支出',12.34,'','','商店','测试','是']
def xlsx(rows):
    book=Workbook();sheet=book.active;sheet.title='记账数据';sheet.append(HEADERS)
    for row in rows: sheet.append(row)
    out=io.BytesIO();book.save(out);book.close();return out.getvalue()

class ImportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='firefly-v101-');self.old=settings._CONFIG_CACHE
        settings._CONFIG_CACHE=Config(database={'url':'sqlite+aiosqlite:///'+self.tmp.name+'/test.db'},auth={'enabled':False},llm={'enabled':False},push={'enabled':True},anomaly={'enabled':False})
        await db.dispose_db();await db.init_db()
        app=FastAPI();app.include_router(api_router,prefix='/api')
        self.client=httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://testserver',trust_env=False)
        self.account=(await self.client.post('/api/accounts',json={'name':'测试钱包','initial_balance':1000})).json()['data']['id']
    async def asyncTearDown(self):
        await self.client.aclose();await db.dispose_db();settings._CONFIG_CACHE=self.old;self.tmp.cleanup()
    async def post(self,rows,**data):
        return (await self.client.post('/api/import/excel',files={'file':('test.xlsx',xlsx(rows))},data={'account_id':str(self.account),**data})).json()
    async def count(self):
        return len((await self.client.get('/api/transactions')).json()['data'])
    async def test_template_is_blank_and_has_rules(self):
        r=await self.client.get('/api/import/excel/template');self.assertEqual(r.status_code,200)
        book=load_workbook(io.BytesIO(r.content));self.assertEqual(book.sheetnames,['记账数据','填写说明']);self.assertIsNone(book['记账数据']['A2'].value);self.assertEqual(book['记账数据'].freeze_panes,'A2');book.close()
        with self.assertRaisesRegex(ValueError,'为空'):excel_import.parse_excel(r.content)
    async def test_preview_never_writes_or_notifies(self):
        with patch.object(excel_import,'notify_import',new_callable=AsyncMock) as notify:
            result=await self.post([ROW],dry_run='true');self.assertTrue(result['success'],result);self.assertTrue(result['data']['dry_run']);self.assertEqual(await self.count(),0);notify.assert_not_called()
    async def test_import_balance_invoice_duplicate_and_notification(self):
        rows=[ROW,['2026-10-17 13:30','收入',20,'','','客户','收入','否']]
        with patch.object(bot,'send_text',new_callable=AsyncMock,return_value=True) as send:
            first=await self.post(rows);self.assertTrue(first['success'],first);self.assertEqual(first['data']['imported_count'],2);self.assertTrue(first['data']['notification']['sent']);send.assert_awaited_once()
            again=await self.post(rows);self.assertEqual(again['data']['duplicate_count'],2);self.assertEqual(send.await_count,1)
        records=(await self.client.get('/api/transactions')).json()['data'];self.assertEqual(len(records),2);self.assertTrue(next(t for t in records if t['direction']=='expense')['has_invoice']);self.assertTrue(all(t['source']=='excel' for t in records))
        accounts=(await self.client.get('/api/accounts')).json()['data'];self.assertAlmostEqual(next(a for a in accounts if a['id']==self.account)['balance'],1007.66)
    async def test_invalid_row_rejects_whole_batch(self):
        result=await self.post([ROW,['2026-10-17','随便',-1,'','','','错误行','否']]);self.assertFalse(result['success']);self.assertEqual(result['detail']['errors'][0]['row'],3);self.assertEqual(await self.count(),0)
    async def test_bad_dates_precision_formula_and_unknown_account(self):
        for date_value,amount in [('错误日期',1),('2026-10-17',1.001),('2026-10-17','=1+1')]:
            result=await self.post([[date_value,'支出',amount,'','','','','']]);self.assertFalse(result['success'],result)
        result=await self.post([['2026-10-17','收入',1,'','不存在账本','','','']]);self.assertFalse(result['success']);self.assertEqual(await self.count(),0)
    async def test_push_failure_does_not_undo_import(self):
        with patch.object(bot,'send_text',new_callable=AsyncMock,return_value=False):
            result=await self.post([ROW]);self.assertTrue(result['success'],result);self.assertFalse(result['data']['notification']['sent']);self.assertEqual(await self.count(),1)
    async def test_cross_account_records_are_not_duplicates(self):
        other=(await self.client.post('/api/accounts',json={'name':'另一钱包'})).json()['data']['id']
        with patch.object(bot,'send_text',new_callable=AsyncMock,return_value=True):
            await self.post([ROW]);result=await self.post([ROW],account_id=str(other));self.assertEqual(result['data']['imported_count'],1)
    async def test_commit_failure_never_notifies(self):
        from sqlmodel.ext.asyncio.session import AsyncSession
        original=AsyncSession.commit
        calls=0
        async def fail_once(session):
            nonlocal calls
            calls+=1
            if calls==2: raise RuntimeError('commit failure')
            return await original(session)
        with patch.object(AsyncSession,'commit',new=fail_once),patch.object(excel_import,'notify_import',new_callable=AsyncMock) as notify:
            result=await self.post([ROW]);self.assertFalse(result['success']);notify.assert_not_called()
        self.assertEqual(await self.count(),0)
    async def test_version_and_update_request_guards(self):
        r=await self.client.get('/api/system/version');self.assertEqual(r.json()['data']['running_version'],__version__)
        self.assertEqual((await self.client.post('/api/system/update')).status_code,400)
        self.assertEqual((await self.client.post('/api/system/update',headers={'X-Firefly-Update':'1','Origin':'https://other.example'})).status_code,403)
    async def test_default_text_destination_and_cq_escaping(self):
        settings.get_config().push.target_type='private';settings.get_config().push.target_id='12345'
        fake=type('Bot',(),{'send_private_msg':AsyncMock(),'send_group_msg':AsyncMock()})()
        with patch.object(bot,'_pick_bot',return_value=fake):self.assertTrue(await bot.send_text('[CQ:at,qq=all] literal'))
        fake.send_private_msg.assert_awaited_once();fake.send_group_msg.assert_not_called();self.assertEqual(fake.send_private_msg.call_args.kwargs['message'][0].type,'text')
    async def test_report_broadcast_target_override(self):
        settings.get_config().push.target_type='private';settings.get_config().push.target_id='12345'
        fake=type('Bot',(),{'send_private_msg':AsyncMock(),'send_group_msg':AsyncMock()})()
        with patch.object(bot,'_pick_bot',return_value=fake):
            result=await bot.broadcast_summary('report',target_type='group',target_id='67890')
        self.assertTrue(result['text_sent']);fake.send_private_msg.assert_not_called()
        self.assertEqual(fake.send_group_msg.call_args.kwargs['group_id'],67890)
    async def test_health_closes_session_and_reports_version(self):
        for _ in range(3):
            payload=(await self.client.get('/api/health')).json()['data']
            self.assertTrue(payload['database']);self.assertEqual(payload['version'],__version__)

class UpdaterTests(unittest.IsolatedAsyncioTestCase):
    async def test_fast_forward_preserves_config_and_stops_for_dirty_source(self):
        with tempfile.TemporaryDirectory(prefix='firefly-update-') as directory:
            base=Path(directory);origin=base/'origin.git';author=base/'author';local=base/'local'
            def git(cwd,*args):
                return subprocess.check_output(['git','-c','user.name=Test','-c','user.email=test@example.com',*args],cwd=cwd,stderr=subprocess.DEVNULL).decode().strip()
            git(base,'init','--bare',str(origin));git(base,'clone',str(origin),str(author));git(author,'switch','-c','main')
            (author/'firefly').mkdir();(author/'firefly/__init__.py').write_text('__version__ = "1.0.0"\n');(author/'config.yaml').write_text('app:\n  version: 1.0.0\npush:\n  target_id: ""\n');(author/'.gitignore').write_text('data/\n')
            git(author,'add','.');git(author,'commit','-m','initial');git(author,'push','origin','main');git(base,'clone','--branch','main',str(origin),str(local))
            (local/'config.yaml').write_text('app:\n  version: 1.0.0\npush:\n  target_id: "local-value"\nauth:\n  secret_key: test-secret\n')
            (author/'firefly/__init__.py').write_text('__version__ = "1.0.1"\n');(author/'config.yaml').write_text('app:\n  version: 1.0.1\npush:\n  target_id: ""\nnew_default: true\n');git(author,'add','.');git(author,'commit','-m','release');git(author,'push','origin','main')
            with patch.object(updater,'REPOSITORY',str(origin).removesuffix('.git')):
                result=await updater.pull_update(local,local/'data');self.assertTrue(result['updated']);self.assertTrue(result['restart_required'])
                cfg=yaml.safe_load((local/'config.yaml').read_text());self.assertEqual(cfg['push']['target_id'],'local-value');self.assertEqual(cfg['auth']['secret_key'],'test-secret');self.assertEqual(cfg['app']['version'],'1.0.1');self.assertTrue(cfg['new_default']);self.assertTrue(list((local/'data/updates').glob('*.yaml')))
                (local/'firefly/__init__.py').write_text('user edits')
                info=await updater.version_info(local);self.assertFalse(info['update_supported']);self.assertIn('未提交',info['reason'])

if __name__=='__main__':unittest.main()
