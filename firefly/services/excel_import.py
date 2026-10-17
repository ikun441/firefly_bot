"""Strict XLSX parsing and atomic financial imports. No formulas or LLM guesses."""
from __future__ import annotations
import asyncio
import io
import zipfile
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo
from decimal import Decimal, InvalidOperation
from typing import Any
from openpyxl import load_workbook
from sqlmodel import select
from ..config import get_config
from ..db.models import Account, ImportBatch, TransactionDirection, TransactionSource
from . import ledger

MAX_ROWS = 2000
MAX_AMOUNT = Decimal('100000000')
HEADERS = {
    'trade_time': ('交易日期', '交易时间', '日期', '发生时间'),
    'direction': ('收支类型', '收支', '收/支', '方向'),
    'amount': ('金额', '金额(元)', '金额（元）', '交易金额'),
    'raw_category': ('类别', '分类'),
    'account_name': ('账本名称', '账本', '账户'),
    'counterparty': ('交易对方', '对方', '商户'),
    'memo': ('备注', '摘要', '用途'),
    'has_invoice': ('是否有发票', '发票', '有无发票'),
}

class ExcelValidationError(ValueError):
    def __init__(self, errors: list[dict]):
        self.errors = errors
        super().__init__(f'发现 {len(errors)} 处填写错误，尚未入账，请按行号修正后重试')

def parse_excel(content: bytes) -> list[dict[str, Any]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError('工作簿解压后过大，请拆分文件')
            if any('vbaproject' in item.filename.lower() for item in archive.infolist()):
                raise ValueError('请上传不含宏的 .xlsx 文件')
            if any(b'<mergeCell' in archive.read(item) for item in archive.infolist()
                   if item.filename.startswith('xl/worksheets/') and item.filename.endswith('.xml')):
                raise ValueError('请取消合并单元格后重新上传')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except zipfile.BadZipFile:
        raise ValueError('无法读取工作簿，请使用未加密、无宏的 .xlsx 文件（解压上限 64MB）') from None
    try:
        sheet = book['记账数据'] if '记账数据' in book.sheetnames else book.worksheets[0]
        if sheet.max_column is None or sheet.max_row is None:
            sheet.calculate_dimension(force=True)
        if sheet.max_column > 32 or sheet.max_row > MAX_ROWS + 1:
            raise ValueError(f'每次最多 {MAX_ROWS} 行、32 列，请拆分文件并清除多余格式')
        rows = sheet.iter_rows()
        header = next(rows, ())
        mapping = {}
        for index, cell in enumerate(header):
            value = str(cell.value or '').strip()
            for field, aliases in HEADERS.items():
                if value in aliases:
                    if field in mapping:
                        raise ValueError(f'表头重复：{value}')
                    mapping[field] = index
        if not {'trade_time', 'direction', 'amount'}.issubset(mapping):
            raise ValueError('首行必须包含交易日期、收支类型、金额，请使用通用模板')
        items, errors = [], []
        for row_number, cells in enumerate(rows, 2):
            if all(cell.value in (None, '') for cell in cells):
                continue
            try:
                if any(cell.data_type in {'f', 'e'} for cell in cells):
                    raise ValueError('不支持公式或错误单元格，请粘贴为实际值')
                row = {field: cells[index].value for field, index in mapping.items()}
                value = row.get('trade_time')
                if isinstance(value, datetime):
                    when = value
                elif isinstance(value, date):
                    when = datetime.combine(value, datetime.min.time())
                else:
                    text = str(value or '').strip().replace('/', '-')
                    try: when = datetime.fromisoformat(text.replace('Z', '+00:00'))
                    except ValueError: raise ValueError('交易日期应为 YYYY-MM-DD 或 YYYY-MM-DD HH:MM') from None
                try: amount = Decimal(str(row.get('amount')))
                except InvalidOperation: raise ValueError('金额必须是数字') from None
                if not amount.is_finite() or not 0 < amount <= MAX_AMOUNT:
                    raise ValueError('金额必须大于 0 且不超过 1 亿元')
                if amount != amount.quantize(Decimal('.01')):
                    raise ValueError('金额最多保留两位小数')
                direction = str(row.get('direction') or '').strip().lower()
                if direction not in {'收入', '支出', 'income', 'expense'}:
                    raise ValueError('收支类型只能填写收入或支出')
                invoice = str(row.get('has_invoice') or '').strip().lower()
                if invoice not in {'', '是', '否', '有', '无', 'true', 'false', '1', '0'}:
                    raise ValueError('是否有发票只能填写是或否（留空视为否）')
                item = {field: str(row.get(field) or '').strip() for field in ['raw_category', 'account_name', 'counterparty', 'memo']}
                if len(item['memo']) > 512 or len(item['counterparty']) > 128:
                    raise ValueError('备注最多 512 字，交易对方最多 128 字')
                if when.tzinfo is None: when = when.replace(tzinfo=ZoneInfo(get_config().app.timezone))
                item.update(row=row_number, trade_time=when.astimezone(timezone.utc).isoformat(), amount=float(amount), direction='income' if direction in {'收入','income'} else 'expense', has_invoice=invoice in {'是','有','true','1'})
                items.append(item)
            except (ValueError, InvalidOperation) as exc:
                errors.append({'row': row_number, 'message': str(exc)})
        if errors: raise ExcelValidationError(errors)
        if not items: raise ValueError('记账数据为空，请填写模板后上传；填写说明页不会导入')
        return items
    finally:
        book.close()

async def resolve_items(session, *, user_id: int, items: list[dict], account_id: int | None) -> list[dict]:
    accounts = list((await session.exec(select(Account).where(Account.user_id == user_id))).all())
    by_name = {}
    for account in accounts:
        by_name.setdefault(account.name, []).append(account)
    default = next((a for a in accounts if a.id == account_id), None)
    if account_id is not None and default is None:
        raise ValueError('所选默认账本不存在或不属于当前用户')
    categories = await ledger.list_categories(session, user_id=user_id)
    errors = []
    for item in items:
        named = by_name.get(item['account_name'], []) if item['account_name'] else []
        account = named[0] if len(named) == 1 else default if not item['account_name'] else None
        if account is None:
            errors.append({'row': item['row'], 'message': '账本不存在、名称重复或未选择默认账本，请先在账本管理中创建'})
        else: item['account_id'] = account.id
        direction = TransactionDirection(item['direction'])
        matches = [c for c in categories if c.direction == direction and c.name == item['raw_category']]
        if item['raw_category'] and len(matches) != 1:
            errors.append({'row': item['row'], 'message': '类别不存在、名称重复或与收支类型不匹配，请使用现有类别'})
        else:
            fallback = next((c for c in categories if c.direction == direction and c.name in {'其他','营业收入','其他收入'}), None)
            item['category_id'] = matches[0].id if matches else fallback.id if fallback else None
    if errors: raise ExcelValidationError(errors)
    return items

async def import_excel_items(session, *, user_id: int, items: list[dict], filename: str) -> dict:
    imported, duplicates = [], []
    income = expense = Decimal('0')
    for item in items:
        when = datetime.fromisoformat(item['trade_time'])
        direction = TransactionDirection(item['direction'])
        hits = await ledger.find_duplicates(session, user_id=user_id, amount=item['amount'], trade_time=when, counterparty=item['counterparty'], direction=direction, account_id=item['account_id'])
        if hits:
            duplicates.append({'row':item['row'], 'duplicate_of':[hit.id for hit in hits]})
            continue
        txn = await ledger.create_transaction(session, user_id=user_id, amount=item['amount'], direction=direction, account_id=item['account_id'], category_id=item['category_id'], trade_time=when, counterparty=item['counterparty'], memo=item['memo'], source=TransactionSource.EXCEL, is_confirmed=True)
        txn.has_invoice = item['has_invoice']; session.add(txn)
        await session.flush()
        from .anomaly import analyze_transaction
        await analyze_transaction(session, user_id=user_id, txn=txn)
        imported.append({'row':item['row'], 'id':txn.id, 'amount':txn.amount, 'direction':item['direction']})
        if direction == TransactionDirection.INCOME: income += Decimal(str(txn.amount))
        else: expense += Decimal(str(txn.amount))
    result = dict(total=len(items), imported_count=len(imported), duplicate_count=len(duplicates), failed_count=0, imported=imported, duplicates=duplicates, income=float(income), expense=float(expense))
    result['filename'] = filename
    result['message'] = f'成功入账 {len(imported)} 笔，重复跳过 {len(duplicates)} 笔'
    batch = ImportBatch(user_id=user_id, source_type='excel', filename=filename[:255], total_count=len(items), imported_count=len(imported), duplicate_count=len(duplicates), pending_count=0, detail=[{'summary':{k:v for k,v in result.items() if k not in {'imported','duplicates'}}}])
    session.add(batch); await session.flush(); result['batch_id'] = batch.id
    return result

async def notify_import(result: dict) -> dict:
    from ..render.bot import get_bot_status, send_text
    if not result['imported_count']:
        return {'attempted':False, 'sent':False, 'message':'没有新增账目，不重复推送'}
    if not get_config().push.enabled:
        return {'attempted':False, 'sent':False, 'message':'推送已在配置中关闭'}
    text = f"流萤 Excel 导入完成\n批次 #{result['batch_id']}\n文件：{result.get('filename', 'Excel')}\n{result['message']}\n本次收入 ¥{result['income']:.2f}\n本次支出 ¥{result['expense']:.2f}"
    try:
        sent = await asyncio.wait_for(send_text(text), timeout=15)
        return {'attempted':True, 'sent':sent, 'message':'导入结果已推送到协议端' if sent else get_bot_status().get('last_error') or '协议端未连接'}
    except Exception:
        return {'attempted':True, 'sent':False, 'message':'推送未完成，账目已保存；请检查协议端连接'}
